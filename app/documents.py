from pathlib import Path
from datetime import datetime
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
import re

OUT=Path(__file__).parent.parent/"data"/"generated"; OUT.mkdir(parents=True,exist_ok=True)

def clean(s): return re.sub(r"[^A-Za-z0-9_-]+","_",s)[:60]

def generate_cv(candidate,job):
    stem=clean(candidate.name+"_"+job["title"]+"_CV")
    doc=Document(); doc.add_heading(candidate.name,0)
    doc.add_paragraph("SAP MM Consultant | SAP S/4HANA | Procurement | P2P")
    doc.add_heading("Professional Summary",1)
    doc.add_paragraph(f"SAP MM professional with {candidate.experience_years:g} years of experience in procurement, inventory management, P2P and SAP operations.")
    doc.add_heading("Core Skills",1); doc.add_paragraph(", ".join(candidate.skills))
    doc.add_heading("Target Role",1); doc.add_paragraph(job["title"]+" — "+job["company"])
    doc.add_heading("Selected Experience Highlights",1)
    for x in ["50,000 materials migrated","350 purchase orders handled","300+ support tickets resolved","Migration Cockpit and SAP Fiori experience","End-user training and SAP MM documentation"]:
        doc.add_paragraph(x,style="List Bullet")
    path=OUT/(stem+".docx"); doc.save(path)
    pdf=OUT/(stem+".pdf")
    styles=getSampleStyleSheet(); story=[Paragraph(candidate.name,styles["Title"]),Paragraph("SAP MM Consultant | SAP S/4HANA | Procurement | P2P",styles["Heading2"]),Spacer(1,10),Paragraph(f"SAP MM professional with {candidate.experience_years:g} years of experience.",styles["BodyText"]),Spacer(1,10),Paragraph("Skills: "+", ".join(candidate.skills),styles["BodyText"]),Spacer(1,10),Paragraph("Selected Highlights",styles["Heading2"])]
    for x in ["50,000 materials migrated","350 purchase orders handled","300+ support tickets resolved","Migration Cockpit and SAP Fiori experience","End-user training and documentation"]:
        story.append(Paragraph("• "+x,styles["BodyText"]))
    SimpleDocTemplate(str(pdf)).build(story)
    return {"docx":str(path),"pdf":str(pdf)}

def cover_letter(candidate,job):
    return f"""Dear Hiring Team,

I am interested in the {job['title']} opportunity at {job['company']}. I bring {candidate.experience_years:g} years of SAP MM experience across procurement, P2P, inventory management, MRP and SAP S/4HANA-related operations.

My experience includes large-scale material migration, purchase-order processing, ticket resolution, SAP Fiori and end-user support. I would welcome the opportunity to contribute to your team and discuss relocation and work-authorization requirements.

Kind regards,
{candidate.name}
"""
